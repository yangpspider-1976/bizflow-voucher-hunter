"""
Builds the "Voucher Hunt — Admin Guide" Word document.

The document is a deliverable, not source: it is regenerated from here so the
wording stays reviewable in git and cannot drift silently from the product.

    pip install python-docx
    python scripts/build-admin-guide.py

Screenshots are optional. Drop PNGs into docs/images/ using the filenames in
SHOTS below and they are embedded automatically; anything missing leaves a
short placeholder line instead.

The role boundaries stated here are the ones enforced in
src/server/admin-users.ts and visibleSections() in
src/app/dashboard/_components/Sidebar.tsx. They are described, never restated
from memory — change one and change the other.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMAGES = ROOT / "docs" / "images"
OUTPUT = ROOT / "docs" / "Voucher Hunt — Admin Guide.docx"

INK = RGBColor(0x0B, 0x1D, 0x3A)
MUTED = RGBColor(0x5B, 0x66, 0x7A)
PURPLE = RGBColor(0x5C, 0x3D, 0xFF)
DANGER = RGBColor(0xB3, 0x26, 0x1E)

SHOTS = {
    "login": ("staff-00-login.png", "Signing in to the dashboard"),
    "business": ("admin-01-business.png", "Creating a business, with the location picker"),
    "campaign": ("admin-02-campaign.png", "The new campaign form"),
    "slots": ("admin-03-slots.png", "Slots for a campaign"),
    "pool": ("admin-04-pool.png", "A voucher tier, and the slots it is offered at"),
    "requests": ("admin-05-requests.png", "Staff slot requests awaiting review"),
    "team": ("admin-06-team.png", "The Team page"),
    "missions": ("admin-07-missions.png", "The mission builder, with its cost simulation"),
    "settings": ("admin-08-settings.png", "Settings, including the Danger Zone"),
}


def shot(doc, key):
    """Embeds the screenshot if it exists, else leaves a light placeholder."""
    filename, caption = SHOTS[key]
    path = IMAGES / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(4.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        line = doc.add_paragraph(caption)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = line.runs[0]
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        return
    line = doc.add_paragraph(f"[ Screenshot: {caption} — docs/images/{filename} ]")
    run = line.runs[0]
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


def warning(doc, text, bold_lead="Warning. "):
    paragraph = doc.add_paragraph()
    lead = paragraph.add_run(bold_lead)
    lead.bold = True
    lead.font.color.rgb = DANGER
    tail = paragraph.add_run(text)
    tail.font.color.rgb = DANGER
    return paragraph


def step(doc, number, title):
    heading = doc.add_heading(f"Step {number} — {title}", level=2)
    heading.runs[0].font.color.rgb = INK
    return heading


def bullets(doc, lines):
    for line in lines:
        if isinstance(line, tuple):
            lead, rest = line
            bullet = doc.add_paragraph(style="List Bullet")
            run = bullet.add_run(lead)
            run.bold = True
            bullet.add_run(rest)
        else:
            doc.add_paragraph(line, style="List Bullet")


def numbered(doc, lines):
    for line in lines:
        if isinstance(line, tuple):
            lead, rest = line
            item = doc.add_paragraph(style="List Number")
            run = item.add_run(lead)
            run.bold = True
            item.add_run(rest)
        else:
            doc.add_paragraph(line, style="List Number")


def result_line(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("You will know it worked: ")
    run.bold = True
    run.font.color.rgb = MUTED
    tail = paragraph.add_run(text)
    tail.font.color.rgb = MUTED
    return paragraph


def two_column_table(doc, rows, headers=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    if headers:
        cells = table.add_row().cells
        for index, heading in enumerate(headers):
            cells[index].text = heading
            cells[index].paragraphs[0].runs[0].bold = True
    for name, detail in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = detail
        cells[0].paragraphs[0].runs[0].bold = True
    return table


def three_column_table(doc, rows, headers):
    table = doc.add_table(rows=0, cols=3)
    table.style = "Light List Accent 1"
    cells = table.add_row().cells
    for index, heading in enumerate(headers):
        cells[index].text = heading
        cells[index].paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
        cells[0].paragraphs[0].runs[0].bold = True
    return table


def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)

    title = doc.add_heading("Voucher Hunt", level=0)
    title.runs[0].font.color.rgb = PURPLE
    subtitle = doc.add_paragraph("Admin & Super Admin Guide — setting up and running the platform")
    subtitle.runs[0].font.size = Pt(13)
    subtitle.runs[0].font.color.rgb = MUTED

    body(
        doc,
        "This guide is for the people who put businesses and campaigns into the "
        "system, create the accounts everyone else signs in with, and keep an "
        "eye on the whole network. It is written in the order a new operator "
        "actually needs it: roles first, then a campaign from nothing, then the "
        "running of it.",
    )

    doc.add_heading("What this app is", level=1)
    body(
        doc,
        "Customers use an Android app to win a discount voucher from a business's "
        "campaign, book a date and time to use it, and show a QR code at the "
        "counter. Alongside that, a network-wide Loyalty Points balance lets "
        "them earn on every purchase and spend with any partner. Your job is the "
        "dashboard behind both.",
    )

    doc.add_heading("What you need before you start", level=1)
    bullets(
        doc,
        [
            "A desktop or laptop with a browser. The dashboard is built for a wide screen.",
            "Your own admin or super admin account — an email and a password.",
            "For a new business: its name, contact number, category and location.",
            "For a new campaign: the offer, the dates, and a decision about which times each prize should be bookable at.",
        ],
    )

    # ---- Roles ------------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 1 — Who can do what", level=1)
    body(
        doc,
        "There are three kinds of dashboard account. The differences are enforced "
        "by the software, not by convention: a link that is not yours is not in "
        "your sidebar, and the page and the request behind it refuse it too.",
    )
    three_column_table(
        doc,
        [
            (
                "Super admin",
                "Everything.",
                "The only role that can reach Team and Settings — that is, create accounts and wipe data.",
            ),
            (
                "Admin",
                "Everything except Team and Settings.",
                "Runs businesses, campaigns, slots, tiers and the network. Cannot create or change accounts.",
            ),
            (
                "Staff",
                "One business, and only the pages that business needs.",
                "The counter and reporting for their own business. Proposes slots and tiers rather than creating them.",
            ),
        ],
        headers=["Role", "Reach", "In practice"],
    )
    body(
        doc,
        "An admin who could create accounts could create itself a super admin, "
        "which would make the boundary decorative. So account management sits "
        "one level up.",
        bold_lead="Why admins cannot create accounts. ",
    )

    doc.add_heading("Guardrails you cannot argue with", level=2)
    bullets(
        doc,
        [
            ("You cannot change your own role or status. ", "Someone else with the rights has to do it. This stops an account locking itself out or promoting itself."),
            ("The last active super admin is protected. ", "They cannot be demoted, deactivated or deleted. Create the replacement first, then change the original."),
            ("A staff account is scoped to its business. ", "Every page it can reach is filtered to that business, so it cannot see another partner's figures."),
        ],
    )

    # ---- Part 2 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 2 — Signing in", level=1)
    numbered(
        doc,
        [
            "Open the dashboard address in your browser.",
            "Enter your Email address and Password.",
            "Select Sign in to Dashboard.",
        ],
    )
    result_line(doc, "the dashboard opens, with your role shown in the account menu at the foot of the sidebar.")
    body(
        doc,
        "The sidebar is grouped into Overview, Manage, Customers and Operations. "
        "If a group is empty for your role it is not shown at all, so you will "
        "never see a heading with nothing under it.",
    )
    shot(doc, "login")

    # ---- Part 3 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 3 — Setting up a campaign from nothing", level=1)
    body(
        doc,
        "Four things, in this order: a business, a campaign, the time slots, then "
        "the voucher tiers. The last two depend on each other, which is why they "
        "come last and together.",
    )

    step(doc, 1, "Create the business")
    numbered(
        doc,
        [
            "Open Businesses and choose to add a new one.",
            "Enter the Business name and Contact number.",
            "Choose the Category.",
            "Set the Location using the picker, so the business appears correctly on the map and in customers' search results.",
            "Save.",
        ],
    )
    result_line(doc, "the business appears in the Businesses list and can be picked on a campaign.")
    shot(doc, "business")

    step(doc, 2, "Create the campaign")
    numbered(
        doc,
        [
            "Open Campaigns and choose to add a new one.",
            "Pick the Business it belongs to.",
            "Enter the Campaign Title. The address customers reach it at is made from the title automatically — there is nothing to type, but a title of only symbols will be refused because it leaves nothing to make an address from.",
            "Write the Offer Message — the line that sells it on the campaign card.",
            "Add a Campaign Image.",
            "Choose the Mode: Restaurant, Online Shop, Beauty, Pet, Retail or Other. This changes the wording customers see and whether they are booking a visit or a window to shop in.",
            "Set the Location, and the Shop URL for an online shop.",
            "Set the Start Date and End Date.",
            "Set Base Attempts — how many spins each customer gets.",
            "Set the Referral Daily Limit — how many extra spins a customer can earn from referrals in a day.",
            "Set the Candidate Timeout (minutes) — how long a customer has to finish booking a prize they have drawn before it goes back in the pool.",
            "Write the Terms.",
            "Save.",
        ],
    )
    result_line(doc, "the campaign appears under Campaigns and can have slots added to it.")
    body(
        doc,
        "On the campaign afterwards you can turn rescheduling on or off. With it "
        "on, staff can move a customer's booking to another slot from the "
        "Scan & Redeem screen; with it off, a voucher is only good in the slot it "
        "was booked at.",
        bold_lead="Rescheduling. ",
    )
    shot(doc, "campaign")

    step(doc, 3, "Open the time slots")
    numbered(
        doc,
        [
            "Open Slots and select the campaign.",
            "Enter the Date, Start Time and End Time.",
            "Set the Total Capacity — how many bookings that window can take.",
            "Save, and repeat for each window.",
        ],
    )
    body(
        doc,
        "A slot can be active, sold out, closed or paused. It goes sold out on "
        "its own when capacity runs out; closing or pausing is yours to do when "
        "something changes at the business.",
    )
    result_line(doc, "the slots are listed against the campaign, with their capacity.")
    shot(doc, "slots")

    step(doc, 4, "Define the voucher tiers")
    body(
        doc,
        "A tier is one kind of prize: what it is worth, how many exist, how often "
        "it is won, and when it can be used.",
    )
    numbered(
        doc,
        [
            "Open Vouchers and select the campaign.",
            "Choose the Benefit Type: a percentage discount, a fixed amount, a free item or free shipping.",
            "Enter the Benefit Value. What this field wants depends on the type — a number for a discount, the name of the item for a free item.",
            "Write the Display Label customers see, for example 20% OFF.",
            "Set the Total Quantity — how many of this prize exist in the whole campaign.",
            "Choose the Rarity: Standard, Rare, Epic or Legendary.",
            "Add a Minimum Spend (optional) and any restriction the offer needs.",
            "Tick the date and time slots this tier is bookable at. At least one is required.",
            "Save.",
        ],
    )
    result_line(doc, "the tier is listed against the campaign and can be drawn by customers.")
    shot(doc, "pool")

    doc.add_heading("The rule that decides the odds", level=2)
    body(
        doc,
        "How often a tier is won comes from its Rarity and nothing else. The "
        "slots you tick decide when the prize can be booked, never how likely it "
        "is. The two are independent, and treating the slot list as an odds "
        "control is the most common mistake made on this screen.",
    )
    three_column_table(
        doc,
        [
            ("Standard", "50", "The everyday prize."),
            ("Rare", "15", "Roughly a third as often as Standard."),
            ("Epic", "5", "Ten times rarer than Standard."),
            ("Legendary", "1", "Fifty times rarer than Standard. The top prize."),
        ],
        headers=["Rarity", "Relative weight", "How often it comes up"],
    )
    body(
        doc,
        "These are ratios between the tiers in a campaign, not percentages. What "
        "a customer actually sees also depends on which tiers still have stock "
        "and bookable slots.",
    )

    doc.add_heading("Designing a campaign that works", level=2)
    bullets(
        doc,
        [
            ("Put the deep discounts on the quiet hours. ", "Tick a Legendary tier only at the slots the business wants to fill. That is the whole point of drawing the prize before the booking."),
            ("Match quantity to capacity. ", "A tier with more vouchers than the slots can seat will frustrate customers who win it and find nothing bookable."),
            ("Add slots before tiers. ", "A tier cannot be saved without at least one slot to attach it to."),
        ],
    )

    # ---- Part 4 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 4 — Approving what staff ask for", level=1)
    body(
        doc,
        "Staff can fill in the same slot and tier forms you can, but what they "
        "submit becomes a request rather than a live change. Nothing exists until "
        "you approve it.",
    )
    numbered(
        doc,
        [
            "Open Slots for slot requests, or Vouchers for voucher tier requests.",
            "Find the Staff Slot Requests or Staff Voucher Tier Requests panel.",
            "Read the row: when it was submitted, who submitted it, and what it would create.",
            "Approve or reject it.",
        ],
    )
    result_line(doc, "the row's status changes from Pending to Approved or Rejected, and records who reviewed it.")
    body(
        doc,
        "Approving creates the slot or tier there and then, with the values in "
        "the request. Rejecting creates nothing. Either way the row stays for "
        "reference, so there is a record of what was asked for and what was "
        "decided.",
    )
    shot(doc, "requests")

    # ---- Part 5 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 5 — Creating accounts", level=1)
    body(
        doc,
        "Super admin only. Everyone who signs in to the dashboard — your own "
        "colleagues and every partner's counter staff — gets an account here. "
        "There is no shared login and no PIN.",
    )
    numbered(
        doc,
        [
            "Open Team.",
            "Add a member and enter their Full name, Email and Password.",
            "Choose the Role: super admin, admin or staff.",
            "For a staff account, choose the Business it is scoped to. This is what limits everything they can see.",
            "Set the Status to active.",
            "Save, and give the person their email and password.",
        ],
    )
    result_line(doc, "the account appears in the Team list and can sign in immediately.")
    body(
        doc,
        "Set their Status to inactive rather than deleting the account. The "
        "history of what they approved and redeemed stays intact, and the sign-in "
        "stops working straight away.",
        bold_lead="When someone leaves. ",
    )
    shot(doc, "team")

    # ---- Part 6 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 6 — Running the network", level=1)
    body(doc, "The day-to-day reporting, and what each page is actually for.")
    two_column_table(
        doc,
        [
            ("Dashboard", "The overview, plus Voucher Benefit Distribution — which tiers are really being won — and the hunt logs."),
            ("Users", "Every customer. Open one for their full history: hunts, vouchers, points."),
            ("Vouchers", "Tiers and the vouchers issued from them, across campaigns."),
            ("Slots", "Capacity and how much of it has gone."),
            ("Transactions", "Checkout history across the network."),
            ("Loyalty Points", "Recent Loyalty Points, and LP Usage & Partner Settlement."),
            ("LP Billing", "Where the Loyalty Points went, Month end, Monthly statements, and the Deposit ledger."),
        ],
        headers=["Page", "What it is for"],
    )

    doc.add_heading("The money, in four lines", level=2)
    two_column_table(
        doc,
        [
            ("A partner awards points", "They owe the network the value of what they issued."),
            ("A customer spends points at a partner", "The network owes the partner 90% of the LP spent; 10% is the service fee."),
            ("Month end", "The two are netted against each other, against the partner's deposit."),
            ("Payout", "The previous month is processed during days 1–7 of the following month."),
        ],
        headers=["Event", "Effect"],
    )
    body(
        doc,
        "Balances are held to exact hundredths of a point, so the 5% earn is "
        "never a rounded guess. The fee is rounded down, and the partner's payout "
        "is always the points spent minus that fee.",
    )

    doc.add_heading("Exports", level=2)
    body(
        doc,
        "Campaign data can be exported as CSV for reporting outside the "
        "dashboard. You can also import a list of already-used codes — a shop's "
        "own used-codes export, for instance — to mark those vouchers redeemed "
        "in bulk rather than one at a time.",
    )

    # ---- Part 7 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 7 — Levels and missions", level=1)
    body(
        doc,
        "This is the part of the product that brings customers back between "
        "campaigns. Levels & Missions is the operator's view of it, with four "
        "pages behind it.",
    )

    doc.add_heading("Levels & Missions", level=2)
    body(
        doc,
        "The control room. It shows the current Economy version and Level ladder "
        "version — the numbers behind experience, rewards and what each level "
        "costs — along with the mission and achievement catalogues. These are "
        "versioned settings you change here, without waiting for a release. Any "
        "events that could not be processed are surfaced at the top so they do "
        "not sit unnoticed.",
    )

    doc.add_heading("Missions", level=2)
    body(
        doc,
        "Where missions are written and approved. A new mission is priced before "
        "it runs: the builder simulates what it would cost if every eligible "
        "customer completed it, so a mission cannot go live on an unexamined "
        "budget. Missions submitted by partners queue here for approval, and the "
        "page leads with how many are waiting.",
    )
    shot(doc, "missions")

    doc.add_heading("Evidence", level=2)
    body(
        doc,
        "Missions that ask for proof — usually a photo of a receipt — land here. "
        "Until someone decides, the customer sees Waiting for review and the "
        "reward is unpaid, so this queue is worth clearing daily. Photos are kept "
        "only while the review needs them and are deleted after 90 days.",
    )

    doc.add_heading("Abuse", level=2)
    body(
        doc,
        "Automated checks run nightly and hold rewards that look wrong rather "
        "than paying them out. The page leads with how many rewards are waiting "
        "to be paid. Release the ones that are fine; the hold is graduated, so a "
        "customer is not cut off by a single odd-looking day.",
    )

    doc.add_heading("Analytics", level=2)
    body(
        doc,
        "The KPI view: Missions, Levels, Economy, Vouchers, and vouchers by "
        "level. This is where you find out whether a change to the economy did "
        "what you wanted.",
    )

    # ---- Part 8 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 8 — Settings and the Danger Zone", level=1)
    body(
        doc,
        "Super admin only. Settings holds platform-level configuration and "
        "maintenance actions.",
    )

    doc.add_heading("SMS delivery", level=2)
    body(
        doc,
        "Controls whether sign-in codes and booking confirmations go out as real "
        "text messages. Customers cannot sign in without receiving a code, so "
        "treat this as a live switch on the front door.",
    )

    doc.add_heading("Danger Zone", level=2)
    warning(
        doc,
        "Both actions below destroy data for everyone, immediately and "
        "permanently. Every customer's vouchers, bookings, Loyalty Points and "
        "history are gone, and every customer and dashboard user is signed out. "
        "There is no undo and no backup taken for you.",
    )
    two_column_table(
        doc,
        [
            ("Reset & Reseed Data", "Empties everything, then recreates the demo data. Type RESET to confirm. For a demo environment being put back to a known state."),
            ("Wipe Data (No Reseed)", "Empties everything and leaves it empty. Type WIPE to confirm."),
        ],
        headers=["Action", "What it does"],
    )
    body(
        doc,
        "The typed confirmation is the only thing standing between a misplaced "
        "click and the whole network. If you are not certain which environment "
        "this dashboard is pointed at, stop and find out first.",
    )
    shot(doc, "settings")

    # ---- Troubleshooting --------------------------------------------------
    doc.add_page_break()
    doc.add_heading("If something goes wrong", level=1)
    two_column_table(
        doc,
        [
            (
                "A campaign is not showing in the app",
                "Check the start and end dates include today, that it has at least one active slot, and that at least one tier is attached to a slot with stock left. A campaign with no bookable prize has nothing to show.",
            ),
            (
                "Customers win a prize but find no times",
                "The tier is attached to slots that are full, closed or past. Attach it to more slots, or open new ones.",
            ),
            (
                "A slot shows sold out with stock left in the tier",
                "Slot capacity and tier quantity are separate limits. The slot is full; the prize is not gone. Open more capacity.",
            ),
            (
                "A tier is never won",
                "Check its rarity. Legendary is fifty times rarer than Standard by design. Also check it still has quantity and a bookable slot.",
            ),
            (
                "A staff member cannot see a page",
                "Staff accounts do not have campaigns, businesses, Team or Settings. If they need a slot or tier, they submit a request and you approve it.",
            ),
            (
                "A staff request cannot be approved",
                "Its values are validated when you approve, not only when it is submitted. A slot whose date has passed, or a tier whose benefit value does not suit its type, is refused with the reason.",
            ),
            (
                "A voucher will not validate at the counter",
                "Find it under Vouchers. Already redeemed, outside its booked slot, or marked no-show all refuse. The customer's own history under Users usually settles it.",
            ),
            (
                "SMS is not arriving",
                "Check SMS delivery under Settings. If it is on and codes still do not arrive, it is a delivery problem outside the dashboard — raise it with whoever runs the SMS service.",
            ),
            (
                "Rewards are not being paid",
                "Check Abuse for held rewards, and Evidence for reviews nobody has decided. Both hold a reward without failing it.",
            ),
            (
                "I cannot change my own role",
                "You cannot, by design. Ask another super admin.",
            ),
        ],
        headers=["Problem", "What to check"],
    )

    # ---- Glossary ---------------------------------------------------------
    doc.add_heading("Words used in this guide", level=1)
    two_column_table(
        doc,
        [
            ("Business", "One partner. A campaign belongs to a business; a staff account is scoped to one."),
            ("Campaign", "A business's offer, running between two dates."),
            ("Mode", "What kind of campaign it is — Restaurant, Online Shop, Beauty, Pet, Retail, Other. Changes the customer wording."),
            ("Slot", "A date and time window with a fixed capacity."),
            ("Voucher pool / tier", "One kind of prize: benefit, quantity, rarity, and the slots it is bookable at."),
            ("Rarity", "Standard, Rare, Epic or Legendary. The only thing that sets a tier's odds."),
            ("Base attempts", "How many spins each customer gets in a campaign."),
            ("Candidate timeout", "How long a drawn prize is held while the customer finishes booking."),
            ("Change request", "A slot or tier a staff account proposed, pending your approval."),
            ("LP (Loyalty Points)", "The network-wide balance customers earn and spend."),
            ("XP", "Experience, which sets a customer's level. Never spendable."),
            ("Settlement", "The monthly netting of what each partner owes and is owed."),
        ],
        headers=["Term", "What it means"],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    embedded = sum(1 for key in SHOTS if (IMAGES / SHOTS[key][0]).exists())
    print(f"Wrote {OUTPUT.name} ({embedded}/{len(SHOTS)} screenshots embedded)")


if __name__ == "__main__":
    build()

"""
Builds the "Voucher Hunt — Staff Guide" Word document.

The document is a deliverable, not source: it is regenerated from here so the
wording stays reviewable in git and cannot drift silently from the product.

    pip install python-docx
    python scripts/build-staff-guide.py

Screenshots are optional. Drop PNGs into docs/images/ using the filenames in
SHOTS below and they are embedded automatically; anything missing leaves a
short placeholder line instead.

What a staff account can see is not a matter of opinion: it is the staffLabels
set in src/app/dashboard/_components/Sidebar.tsx. Change one and change the
other.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMAGES = ROOT / "docs" / "images"
OUTPUT = ROOT / "docs" / "Voucher Hunt — Staff Guide.docx"

INK = RGBColor(0x0B, 0x1D, 0x3A)
MUTED = RGBColor(0x5B, 0x66, 0x7A)
PURPLE = RGBColor(0x5C, 0x3D, 0xFF)

SHOTS = {
    "login": ("staff-00-login.png", "Signing in to the dashboard"),
    "validate": ("staff-01-validate.png", "Scan & Redeem, with a code entered"),
    "awarded": ("staff-02-awarded.png", "The 5% confirmed at the counter"),
    "result": ("staff-03-result.png", "A validation result, before marking it used"),
    "dashboard": ("staff-04-dashboard.png", "The dashboard overview for a business"),
    "billing": ("staff-05-billing.png", "LP Billing, showing the monthly statement"),
    "requests": ("staff-06-requests.png", "A submitted slot request, awaiting approval"),
}


def shot(doc, key):
    """Embeds the screenshot if it exists, else leaves a light placeholder."""
    filename, caption = SHOTS[key]
    path = IMAGES / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(4.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        line = doc.add_paragraph(caption)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = line.runs[0]
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        return
    line = doc.add_paragraph(f"[ Screenshot: {caption} — docs/images/{filename} ]")
    run = line.runs[0]
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


def step(doc, number, title):
    heading = doc.add_heading(f"Step {number} — {title}", level=2)
    heading.runs[0].font.color.rgb = INK
    return heading


def bullets(doc, lines):
    for line in lines:
        if isinstance(line, tuple):
            lead, rest = line
            bullet = doc.add_paragraph(style="List Bullet")
            run = bullet.add_run(lead)
            run.bold = True
            bullet.add_run(rest)
        else:
            doc.add_paragraph(line, style="List Bullet")


def numbered(doc, lines):
    for line in lines:
        if isinstance(line, tuple):
            lead, rest = line
            item = doc.add_paragraph(style="List Number")
            run = item.add_run(lead)
            run.bold = True
            item.add_run(rest)
        else:
            doc.add_paragraph(line, style="List Number")


def result_line(doc, text):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("You will know it worked: ")
    run.bold = True
    run.font.color.rgb = MUTED
    tail = paragraph.add_run(text)
    tail.font.color.rgb = MUTED
    return paragraph


def two_column_table(doc, rows, headers=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    if headers:
        cells = table.add_row().cells
        for index, heading in enumerate(headers):
            cells[index].text = heading
            cells[index].paragraphs[0].runs[0].bold = True
    for name, detail in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = detail
        cells[0].paragraphs[0].runs[0].bold = True
    return table


def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)

    title = doc.add_heading("Voucher Hunt", level=0)
    title.runs[0].font.color.rgb = PURPLE
    subtitle = doc.add_paragraph("Staff Guide — running your business on the dashboard")
    subtitle.runs[0].font.size = Pt(13)
    subtitle.runs[0].font.color.rgb = MUTED

    body(
        doc,
        "This guide is for business owners and the people on their counter. The "
        "job at the till is two screens and about thirty seconds, and that part "
        "is covered first and in the most detail. Everything after it is the "
        "reporting you will look at weekly rather than hourly.",
    )

    doc.add_heading("What this app is", level=1)
    body(
        doc,
        "Customers use an Android app to win a discount voucher from your "
        "campaign and book a specific date and time to come and use it. They "
        "turn up inside that slot and show a QR code. You scan it, confirm it, "
        "and apply the discount on your own till as normal. Separately, "
        "customers collect Loyalty Points — LP — worth 5% of what they spend "
        "with you, and can come back and spend those points across the network.",
    )

    doc.add_heading("What you need before you start", level=1)
    bullets(
        doc,
        [
            "A desktop, laptop or tablet with a web browser. There is nothing to install.",
            "The dashboard address for your business, and the email and password an admin created for you. There is no PIN and no shared login — the account is yours.",
            "A webcam if you want to scan QR codes directly. Without one you can still upload a photo of the code or type it in, and the scan button is simply hidden.",
        ],
    )

    # ---- Part 1 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 1 — Signing in and finding your way", level=1)

    step(doc, 1, "Sign in")
    numbered(
        doc,
        [
            "Open the dashboard address in your browser.",
            "Enter your Email address and Password.",
            "Select Sign in to Dashboard.",
        ],
    )
    result_line(doc, "the dashboard opens with your business name in the sidebar.")
    shot(doc, "login")

    doc.add_heading("What you can see", level=2)
    body(
        doc,
        "A staff account sees a deliberately narrow dashboard, and everything on "
        "it is scoped to your own business — your customers, your history, your "
        "money. You will not see other businesses' figures anywhere.",
    )
    two_column_table(
        doc,
        [
            ("Dashboard", "The overview: how the campaign is going, at a glance."),
            ("Slots", "The date and time windows customers can book, and their capacity."),
            ("Vouchers", "The benefit tiers on offer and the vouchers issued from them."),
            ("Users", "The customers who took part, and one customer's history."),
            ("Loyalty Points", "Points earned and spent with you."),
            ("Transactions", "Your checkout's own history."),
            ("LP Billing", "What you owe and what you are owed, and the monthly statement."),
            ("Levels & Missions", "Missions you run, and the customer evidence waiting for review."),
            ("Scan & Redeem", "The counter screen. The one you will use every day."),
        ],
        headers=["Sidebar link", "What it is for"],
    )
    body(
        doc,
        "Creating businesses and campaigns, managing team accounts, and platform "
        "Settings are not available to a staff account. Slots and voucher tiers "
        "you can ask for — see Part 4.",
        bold_lead="What is not there. ",
    )

    # ---- Part 2 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 2 — Scan & Redeem: the counter job", level=1)
    body(
        doc,
        "This is the screen to keep open during service. It handles every kind "
        "of code a customer can present, and you do not need to know which one "
        "you are holding — put it in the same box either way.",
    )

    doc.add_heading("Accepting a voucher", level=2)
    step(doc, 1, "Take the code")
    numbered(
        doc,
        [
            "Open Scan & Redeem.",
            "Select Scan QR and hold the customer's screen to the camera. If there is no camera, select Upload QR Image, or type the code into the Voucher Code or QR Token box — they look like BF20-15MAY-12PM-X7A8.",
            "Select Validate voucher.",
        ],
    )
    result_line(doc, "a Validation Result panel appears below the box.")
    shot(doc, "validate")

    step(doc, 2, "Check the result")
    body(
        doc,
        "The result panel tells you everything you need to decide, without "
        "asking the customer anything:",
    )
    two_column_table(
        doc,
        [
            ("Customer", "Who the booking is under."),
            ("Benefit", "The discount to apply."),
            ("Selected Slot", "The date and time they booked. Check it against the clock."),
            ("Expires", "When the voucher stops being valid — the end of that slot."),
            ("Minimum Spend", "A floor on the bill, where the tier sets one."),
            ("Remaining", "What is left on the voucher, where it can be part-used."),
        ],
        headers=["The panel shows", "What to do with it"],
    )
    body(
        doc,
        "A voucher that is already used, expired or not genuine is refused here, "
        "with the reason. There is nothing to work out — if the panel refuses "
        "it, it cannot be accepted.",
        bold_lead="If it is not valid. ",
    )
    shot(doc, "result")

    step(doc, 3, "Apply the discount and mark it used")
    numbered(
        doc,
        [
            "Apply the discount on your own till as you normally would.",
            "Enter what the customer actually paid in Amount paid (₱).",
            "Add an Internal Note (optional) if anything needs recording.",
            "Select Mark as Used.",
        ],
    )
    body(
        doc,
        "The amount paid is what the customer's 5% is calculated from. Mark the "
        "voucher used without it and the voucher is still correctly spent, but "
        "that sale earns the customer no points.",
        bold_lead="Why the amount matters. ",
    )
    result_line(doc, "the screen confirms the redemption and shows the points the customer earned.")
    shot(doc, "awarded")

    doc.add_heading("Giving points after an ordinary purchase", level=2)
    body(
        doc,
        "A customer does not need a voucher to earn points. After any paid "
        "purchase they can show you the wallet QR from the More tab of their app.",
    )
    numbered(
        doc,
        [
            "Take payment as normal.",
            "Scan the customer's wallet QR, or paste the wallet token into the same box.",
            "Enter the Purchase amount.",
            "Confirm.",
        ],
    )
    body(
        doc,
        "The app works out 5% and credits it — ₱500 spent becomes 25 LP — and "
        "tells you what was added so you can tell the customer. The same sale "
        "cannot be counted twice however many times the button is pressed.",
    )
    body(
        doc,
        "Occasionally you will see Loyalty Points held for review instead of a "
        "balance. Nothing is wrong at your end and the sale is fine; the points "
        "are simply waiting on a check before they land.",
        bold_lead="Held for review. ",
    )

    doc.add_heading("Taking Loyalty Points as payment", level=2)
    body(
        doc,
        "Customers can pay with an LP voucher, or collect an item they already "
        "bought with points in the app. Both arrive as a code in the same box.",
    )
    bullets(
        doc,
        [
            ("An item bought in the app. ", "The code names the item and the partner itself. Select Confirm Handover and give them the item."),
            ("A plain LP voucher. ", "It can be spent at any partner, so the screen asks you to pick who is accepting it under Redeeming at, then enter the bill. Select Accept LP Payment."),
            ("Part of a bill is fine. ", "A voucher does not have to cover the whole sale. What is left stays on the voucher for next time."),
        ],
    )
    body(
        doc,
        "If your account is scoped to a single business, Redeeming at is filled "
        "in for you and there is nothing to choose.",
    )

    doc.add_heading("No-shows and rescheduling", level=2)
    bullets(
        doc,
        [
            ("They did not turn up. ", "Validate the voucher and select Mark No-show. The booking is recorded against that customer and the place is given up."),
            ("They need a different time. ", "If the campaign allows it, a Reschedule option appears next to the result with the other slots you could move them to. If the campaign does not allow it, the option is not there — that is the campaign's setting, not a fault."),
        ],
    )

    # ---- Part 3 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 3 — Watching how it is going", level=1)

    doc.add_heading("Dashboard", level=2)
    body(
        doc,
        "The overview page opens on your headline numbers, then two panels worth "
        "knowing: Voucher Benefit Distribution, which shows which tiers are "
        "actually being won, and User Attempts / Voucher Hunt Logs, which shows "
        "the hunts people are running.",
    )
    shot(doc, "dashboard")

    doc.add_heading("Slots", level=2)
    body(
        doc,
        "Every date and time window you have open, with its capacity and how "
        "much of it is gone. A slot fills up and closes itself; you do not have "
        "to watch it. This is also where you ask for new slots — Part 4.",
    )

    doc.add_heading("Vouchers", level=2)
    body(
        doc,
        "The benefit tiers on your campaign and the vouchers issued from them. "
        "Use it to see what is left in a tier before a busy day.",
    )

    doc.add_heading("Users", level=2)
    body(
        doc,
        "The customers who took part. Open one to see their history with you — "
        "useful when someone at the counter has a question about a booking that "
        "is not in front of you.",
    )

    doc.add_heading("Transactions", level=2)
    body(
        doc,
        "Your own checkout's history, line by line. Everything here is your "
        "business and nobody else's.",
    )

    doc.add_heading("Loyalty Points and LP Billing", level=2)
    body(
        doc,
        "Loyalty Points shows Recent Loyalty Points and LP Usage & Partner "
        "Settlement. LP Billing is the money view, and has four parts: Where the "
        "Loyalty Points went, Month end, Monthly statements, and the Deposit "
        "ledger.",
    )
    body(doc, "The arithmetic behind those pages is short:")
    two_column_table(
        doc,
        [
            ("When you award points", "You owe the network their value. 50 LP issued costs you ₱50."),
            ("When a customer spends points with you", "The network owes you. Of the LP spent, you receive 90% — the other 10% is the service fee."),
            ("At month end", "The two are netted against each other, against your deposit."),
            ("When you are paid", "Settlement for the previous month is processed during days 1–7 of the following month."),
        ],
        headers=["Situation", "What happens"],
    )
    shot(doc, "billing")

    doc.add_heading("Levels & Missions", level=2)
    body(
        doc,
        "Missions are small tasks that send customers to you. You can write your "
        "own, and the dashboard prices them before they run so you can see what "
        "a mission would cost if everyone completed it. A mission you submit is "
        "reviewed before it goes live.",
    )
    body(
        doc,
        "Some missions ask the customer for evidence — usually a photo of a "
        "receipt. Those arrive in the evidence queue and wait for a person to "
        "look. Until someone does, the customer sees Waiting for review and the "
        "reward is not paid, so it is worth clearing the queue daily.",
        bold_lead="The evidence queue. ",
    )

    # ---- Part 4 -----------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Part 4 — Asking for a new slot or voucher tier", level=1)
    body(
        doc,
        "You can open new times and propose new benefit tiers, but they do not "
        "go live the moment you save them. What you fill in becomes a request, "
        "and it takes effect only once an admin approves it. This is deliberate: "
        "capacity and discounts are what the campaign costs, so a second pair of "
        "eyes sits in front of them.",
    )

    doc.add_heading("Asking for a new time slot", level=2)
    numbered(
        doc,
        [
            "Open Slots.",
            "Fill in the Date, Start Time, End Time and Total Capacity.",
            "Submit the form.",
            "Find it under Your Slot Requests, where it will show as Pending.",
        ],
    )
    result_line(doc, "the request appears in Your Slot Requests with a Pending badge.")
    shot(doc, "requests")

    doc.add_heading("Asking for a new voucher tier", level=2)
    numbered(
        doc,
        [
            "Open Vouchers.",
            "Choose the Benefit Type — a percentage discount, a fixed amount, a free item or free shipping — and enter the Benefit Value.",
            "Write the Display Label the customer will see, for example 20% OFF.",
            "Set the Total Quantity, and a Minimum Spend (optional) if the offer needs one.",
            "Choose the Rarity, which decides how often the tier is won.",
            "Tick the date and time slots the tier should be bookable at. At least one is required.",
            "Submit, and find it under Your Voucher Tier Requests.",
        ],
    )
    body(
        doc,
        "How often a tier is won comes from its Rarity alone. Offering it at more "
        "slots makes it easier to book, not easier to win.",
        bold_lead="One thing people get wrong. ",
    )

    doc.add_heading("What the three states mean", level=2)
    two_column_table(
        doc,
        [
            ("Pending", "Submitted and waiting. Nothing is live yet."),
            ("Approved", "An admin accepted it. The slot or tier now exists and customers can reach it."),
            ("Rejected", "An admin declined it. Nothing was created. The row stays for reference."),
        ],
        headers=["Status", "What it means"],
    )

    # ---- Troubleshooting --------------------------------------------------
    doc.add_page_break()
    doc.add_heading("If something goes wrong", level=1)
    two_column_table(
        doc,
        [
            (
                "The code is not recognised",
                "Check it was typed in full. If the customer is reading it out, ask for the QR instead. A code that is genuinely invalid, used or expired is refused with the reason.",
            ),
            (
                "There is no Scan QR button",
                "That machine has no camera, so the button is hidden rather than offered and broken. Use Upload QR Image, or type the code.",
            ),
            (
                "The camera will not start",
                "The browser needs permission to use it. Allow it when prompted, or use Upload QR Image instead.",
            ),
            (
                "It says the voucher is already redeemed",
                "It has been used. Vouchers cannot be used twice. Check with the customer whether someone else in the party already presented it.",
            ),
            (
                "The customer is here at the wrong time",
                "The booked slot is on the result panel. If the campaign allows rescheduling, use Reschedule to move them; otherwise the voucher is only valid in its own slot.",
            ),
            (
                "Points were not added to the customer",
                "The 5% comes from Amount paid. If that was left blank, the sale earned nothing. Credit them from the wallet QR instead.",
            ),
            (
                "The screen says points are held for review",
                "The sale is fine and the voucher is accepted. The points are waiting on a routine check.",
            ),
            (
                "My slot request is still Pending",
                "It needs an admin to approve it. Nothing is live until they do.",
            ),
            (
                "A page I expected is not in the sidebar",
                "Staff accounts do not have campaigns, businesses, Team or Settings. Ask an admin for what you need.",
            ),
        ],
        headers=["Problem", "What to do"],
    )

    # ---- Glossary ---------------------------------------------------------
    doc.add_heading("Words used in this guide", level=1)
    two_column_table(
        doc,
        [
            ("Campaign", "Your offer, running between two dates, that customers hunt."),
            ("Slot", "A date and time window customers book, with a fixed capacity."),
            ("Voucher pool / tier", "One kind of prize in a campaign — the benefit, how many exist, and the slots it is bookable at."),
            ("Rarity", "How often a tier is won: Standard, Rare, Epic or Legendary."),
            ("LP (Loyalty Points)", "The network-wide balance customers earn and spend."),
            ("LP voucher", "Points turned into a code, spendable at a partner, and usable against part of a bill."),
            ("Wallet QR", "The code in the customer's More tab that you scan to award the 5%."),
            ("No-show", "A booking the customer did not turn up for, recorded against them."),
            ("Change request", "A slot or tier you proposed, waiting for an admin to approve."),
            ("Settlement", "The monthly netting of what you owe and what you are owed."),
        ],
        headers=["Term", "What it means"],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    embedded = sum(1 for key in SHOTS if (IMAGES / SHOTS[key][0]).exists())
    print(f"Wrote {OUTPUT.name} ({embedded}/{len(SHOTS)} screenshots embedded)")


if __name__ == "__main__":
    build()
